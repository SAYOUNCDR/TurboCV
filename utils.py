import json
import os
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document
from jinja2 import Environment, FileSystemLoader, select_autoescape

try:
    import google.generativeai as genai
except Exception:  # pragma: no cover
    genai = None

try:
    import pdfkit
except Exception:  # pragma: no cover
    pdfkit = None

try:
    from groq import Groq
except Exception:  # pragma: no cover
    Groq = None


BASE_DIR = Path(__file__).resolve().parent
TEMPLATES_DIR = BASE_DIR / "templates"
OUTPUT_DIR = BASE_DIR / "generated"

SUPPORTED_JD_FORMATS = "text message, JPG, PNG, JPEG, WEBP screenshot, or image document"
SUPPORTED_RESUME_FORMATS = "PDF, DOCX, TXT, MD, pasted text, or Notion export saved as text/markdown/PDF"

JD_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
TEXT_EXTENSIONS = {".txt", ".md", ".rtf"}
ACTION_VERBS = {
    "built",
    "created",
    "delivered",
    "designed",
    "drove",
    "improved",
    "implemented",
    "increased",
    "launched",
    "led",
    "managed",
    "optimized",
    "reduced",
    "scaled",
}
COMMON_SKILLS = [
    "python",
    "java",
    "javascript",
    "typescript",
    "react",
    "next.js",
    "node.js",
    "django",
    "flask",
    "fastapi",
    "sql",
    "postgresql",
    "mysql",
    "mongodb",
    "redis",
    "docker",
    "kubernetes",
    "aws",
    "azure",
    "gcp",
    "linux",
    "git",
    "rest",
    "graphql",
    "ci/cd",
    "machine learning",
    "data analysis",
    "nlp",
    "llm",
    "genai",
    "figma",
    "html",
    "css",
]
STOPWORDS = {
    "a",
    "an",
    "and",
    "are",
    "as",
    "at",
    "be",
    "by",
    "for",
    "from",
    "have",
    "in",
    "is",
    "of",
    "on",
    "or",
    "that",
    "the",
    "to",
    "with",
    "will",
    "you",
    "your",
    "their",
    "our",
    "this",
    "role",
    "team",
    "work",
    "years",
    "year",
    "experience",
    "using",
    "strong",
    "skills",
    "skill",
}


def generated_output_dir(kind: str, chat_id: int) -> Path:
    path = OUTPUT_DIR / kind / str(chat_id)
    path.mkdir(parents=True, exist_ok=True)
    return path


def infer_extension(file_name: str, mime_type: str | None) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix:
        return suffix

    mime_map = {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/webp": ".webp",
        "application/pdf": ".pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
        "text/plain": ".txt",
        "text/markdown": ".md",
    }
    return mime_map.get(mime_type or "", ".bin")


def extract_jd_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension in JD_IMAGE_EXTENSIONS:
        return _extract_text_from_image(path)
    if extension == ".pdf":
        return _extract_pdf_text(path)
    if extension == ".docx":
        return _extract_docx_text(path)
    if extension in TEXT_EXTENSIONS:
        return _extract_text_file(path)
    raise ValueError("Unsupported JD format. Please send text or an image screenshot.")


def extract_resume_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension == ".pdf":
        return _extract_pdf_text(path)
    if extension == ".docx":
        return _extract_docx_text(path)
    if extension in TEXT_EXTENSIONS:
        return _extract_text_file(path)
    raise ValueError("Unsupported resume format. Please upload PDF, DOCX, TXT, or MD.")


def analyze_resume_against_jd(jd_text: str, resume_text: str, chat_id: int) -> dict[str, Any]:
    analysis = _analyze_with_ai(jd_text, resume_text) or _analyze_locally(jd_text, resume_text)
    draft_resume = analysis.get("draft_resume") or _build_resume_draft(
        jd_text,
        resume_text,
        analysis,
    )
    analysis["draft_resume"] = draft_resume
    analysis["generated_files"] = [
        str(path) for path in _generate_resume_files(chat_id, analysis)
    ]
    return analysis


def answer_general_query(
    query: str,
    jd_text: str | None = None,
    resume_text: str | None = None,
    last_analysis: dict[str, Any] | None = None,
) -> str:
    response = _answer_general_query_with_ai(query, jd_text, resume_text, last_analysis)
    if response:
        return response
    return _answer_general_query_locally(query, jd_text, resume_text, last_analysis)


def _extract_pdf_text(path: Path) -> str:
    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                parts.append(page_text.strip())
    text = "\n\n".join(parts).strip()
    if not text:
        raise ValueError("No readable text was found in the PDF.")
    return text


def _extract_docx_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("No readable text was found in the DOCX file.")
    return text


def _extract_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = path.read_text(encoding=encoding).strip()
            if text:
                return text
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not read the text file.")


def _extract_text_from_image(path: Path) -> str:
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key or genai is None:
        raise ValueError(
            "JD screenshots need GEMINI_API_KEY configured. You can also paste the JD as text."
        )

    genai.configure(api_key=api_key)
    model = genai.GenerativeModel(os.getenv("GEMINI_VISION_MODEL", "gemini-1.5-flash"))
    uploaded = genai.upload_file(path=str(path))
    prompt = (
        "Extract the complete job description from this image. "
        "Return clean plain text only. Preserve key requirements, tools, and qualifications."
    )
    response = model.generate_content([uploaded, prompt])
    text = getattr(response, "text", "").strip()
    if not text:
        raise ValueError("The screenshot did not return readable JD text.")
    return text


def _answer_general_query_with_ai(
    query: str,
    jd_text: str | None,
    resume_text: str | None,
    last_analysis: dict[str, Any] | None,
) -> str | None:
    groq_key = os.getenv("GROQ_API_KEY")
    if groq_key and Groq is not None:
        try:
            return _chat_with_groq(query, jd_text, resume_text, last_analysis, groq_key)
        except Exception:
            pass

    gemini_key = os.getenv("GEMINI_API_KEY")
    if gemini_key and genai is not None:
        try:
            return _chat_with_gemini(query, jd_text, resume_text, last_analysis, gemini_key)
        except Exception:
            pass

    return None


def _analyze_with_ai(jd_text: str, resume_text: str) -> dict[str, Any] | None:
    groq_key = os.getenv("GROQ_API_KEY")
    if groq_key and Groq is not None:
        try:
            return _analyze_with_groq(jd_text, resume_text, groq_key)
        except Exception:
            pass

    gemini_key = os.getenv("GEMINI_API_KEY")
    if gemini_key and genai is not None:
        try:
            return _analyze_with_gemini(jd_text, resume_text, gemini_key)
        except Exception:
            pass

    return None


def _analyze_with_groq(jd_text: str, resume_text: str, api_key: str) -> dict[str, Any]:
    client = Groq(api_key=api_key)
    prompt = _analysis_prompt(jd_text, resume_text)
    response = client.chat.completions.create(
        model=os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile"),
        temperature=0.2,
        response_format={"type": "json_object"},
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a strict resume evaluator. Return valid JSON only. "
                    "Never invent facts. If information is missing, use placeholders."
                ),
            },
            {"role": "user", "content": prompt},
        ],
    )
    content = response.choices[0].message.content
    return _normalize_ai_payload(_parse_json_payload(content))


def _analyze_with_gemini(jd_text: str, resume_text: str, api_key: str) -> dict[str, Any]:
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel(os.getenv("GEMINI_TEXT_MODEL", "gemini-1.5-flash"))
    response = model.generate_content(_analysis_prompt(jd_text, resume_text))
    content = getattr(response, "text", "")
    return _normalize_ai_payload(_parse_json_payload(content))


def _chat_with_groq(
    query: str,
    jd_text: str | None,
    resume_text: str | None,
    last_analysis: dict[str, Any] | None,
    api_key: str,
) -> str:
    client = Groq(api_key=api_key)
    response = client.chat.completions.create(
        model=os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile"),
        temperature=0.4,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are TurboCV, a helpful career assistant inside a Telegram bot. "
                    "Answer clearly and briefly. If resume or JD context is provided, use it. "
                    "Do not invent personal facts not present in the context."
                ),
            },
            {"role": "user", "content": _general_chat_prompt(query, jd_text, resume_text, last_analysis)},
        ],
    )
    return (response.choices[0].message.content or "").strip()


def _chat_with_gemini(
    query: str,
    jd_text: str | None,
    resume_text: str | None,
    last_analysis: dict[str, Any] | None,
    api_key: str,
) -> str:
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel(os.getenv("GEMINI_TEXT_MODEL", "gemini-1.5-flash"))
    response = model.generate_content(_general_chat_prompt(query, jd_text, resume_text, last_analysis))
    return getattr(response, "text", "").strip()


def _general_chat_prompt(
    query: str,
    jd_text: str | None,
    resume_text: str | None,
    last_analysis: dict[str, Any] | None,
) -> str:
    score = ""
    if last_analysis:
        score = (
            f"Last known analysis score: {last_analysis.get('score', 'unknown')} / 10\n"
            f"Alignment: {last_analysis.get('alignment_label', 'unknown')}\n"
            f"Summary: {last_analysis.get('summary', '')}\n"
        )

    jd_snippet = (jd_text or "")[:1800]
    resume_snippet = (resume_text or "")[:1800]
    return f"""
Answer the user's career or resume question as TurboCV.
Be concise, practical, and supportive.
If useful, refer to the available JD, resume, or prior analysis context.
If context is missing, answer generally and say what extra input would improve the answer.

User question:
{query}

Context:
{score}
Job Description:
{jd_snippet or "Not available"}

Resume:
{resume_snippet or "Not available"}
""".strip()


def _analysis_prompt(jd_text: str, resume_text: str) -> str:
    return f"""
Analyze this resume against this job description and return strict JSON only.

Required JSON shape:
{{
  "score": 0.0,
  "alignment_label": "Strong fit | Moderate fit | Needs improvement",
  "summary": "short paragraph",
  "matched_keywords": ["keyword"],
  "missing_keywords": ["keyword"],
  "suggestions": ["specific suggestion"],
  "draft_resume": {{
    "name": "candidate name or Candidate Name",
    "headline": "target headline",
    "contact": "contact line or placeholder",
    "summary": "tailored summary without inventing facts",
    "skills": ["skill"],
    "experience": ["bullet"],
    "projects": ["bullet"],
    "education": ["bullet"],
    "certifications": ["bullet"],
    "notes": ["placeholder note only when needed"]
  }}
}}

Rules:
- Score must be out of 10.
- Do not invent experience or companies.
- Use placeholders like [Add metric from your real work] only when necessary.
- Suggestions must be practical and action oriented.
- Keep matched_keywords and missing_keywords concise.

Job Description:
{jd_text}

Resume:
{resume_text}
""".strip()


def _parse_json_payload(content: str) -> dict[str, Any]:
    cleaned = content.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`")
        cleaned = cleaned.replace("json", "", 1).strip()

    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, re.DOTALL)
        if not match:
            raise
        return json.loads(match.group(0))


def _normalize_ai_payload(payload: dict[str, Any]) -> dict[str, Any]:
    return {
        "score": round(float(payload.get("score", 5.0)), 1),
        "alignment_label": payload.get("alignment_label", "Needs improvement"),
        "summary": str(payload.get("summary", "Resume analysis completed.")),
        "matched_keywords": _clean_string_list(payload.get("matched_keywords", [])),
        "missing_keywords": _clean_string_list(payload.get("missing_keywords", [])),
        "suggestions": _clean_string_list(payload.get("suggestions", [])),
        "draft_resume": _normalize_draft(payload.get("draft_resume", {})),
    }


def _normalize_draft(payload: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": str(payload.get("name", "Candidate Name")).strip() or "Candidate Name",
        "headline": str(payload.get("headline", "Target Role Resume")).strip() or "Target Role Resume",
        "contact": str(payload.get("contact", "Add email | Add phone | Add LinkedIn")).strip(),
        "summary": str(payload.get("summary", "")).strip(),
        "skills": _clean_string_list(payload.get("skills", [])),
        "experience": _clean_string_list(payload.get("experience", [])),
        "projects": _clean_string_list(payload.get("projects", [])),
        "education": _clean_string_list(payload.get("education", [])),
        "certifications": _clean_string_list(payload.get("certifications", [])),
        "notes": _clean_string_list(payload.get("notes", [])),
    }


def _clean_string_list(values: Any) -> list[str]:
    if not isinstance(values, list):
        return []
    cleaned = []
    for value in values:
        text = str(value).strip()
        if text:
            cleaned.append(text)
    return cleaned


def _analyze_locally(jd_text: str, resume_text: str) -> dict[str, Any]:
    jd_keywords = _extract_keywords(jd_text)
    resume_keywords = set(_extract_keywords(resume_text))
    matched = [keyword for keyword in jd_keywords if keyword in resume_keywords][:12]
    missing = [keyword for keyword in jd_keywords if keyword not in resume_keywords][:12]

    sections = {
        "summary": any(token in resume_text.lower() for token in ("summary", "profile", "objective")),
        "skills": "skills" in resume_text.lower(),
        "experience": any(token in resume_text.lower() for token in ("experience", "employment", "work history")),
        "projects": "project" in resume_text.lower(),
        "education": "education" in resume_text.lower(),
    }
    section_score = sum(sections.values()) / len(sections)
    keyword_ratio = len(matched) / max(len(jd_keywords), 1)
    action_score = 1.0 if any(verb in resume_text.lower() for verb in ACTION_VERBS) else 0.4
    metrics_score = 1.0 if re.search(r"(\d+%|\d+\+|\$\d+)", resume_text) else 0.4
    total = 2.0 + keyword_ratio * 5.0 + section_score * 1.5 + action_score * 0.75 + metrics_score * 0.75
    score = round(max(1.0, min(total, 10.0)), 1)

    if score >= 8:
        label = "Strong fit"
    elif score >= 6:
        label = "Moderate fit"
    else:
        label = "Needs improvement"

    suggestions = []
    if missing:
        suggestions.append(
            f"Reflect verified JD keywords more clearly, especially: {', '.join(missing[:5])}."
        )
    if not sections["skills"]:
        suggestions.append("Add a dedicated skills section aligned with the JD.")
    if metrics_score < 1.0:
        suggestions.append("Add measurable impact, such as percentages, time saved, or revenue influence.")
    if action_score < 1.0:
        suggestions.append("Rewrite bullets with stronger action verbs and outcome-focused language.")
    suggestions.append("Tailor the resume summary so the first lines mirror the target role more closely.")

    summary = (
        f"The resume shows a {label.lower()} against the JD. "
        f"It matches {len(matched)} important keywords and misses {len(missing)} higher-priority JD terms."
    )

    return {
        "score": score,
        "alignment_label": label,
        "summary": summary,
        "matched_keywords": matched,
        "missing_keywords": missing,
        "suggestions": suggestions[:5],
        "draft_resume": _build_resume_draft(
            jd_text,
            resume_text,
            {
                "matched_keywords": matched,
                "missing_keywords": missing,
                "suggestions": suggestions[:5],
            },
        ),
    }


def _extract_keywords(text: str) -> list[str]:
    lowered = text.lower()
    keywords: list[str] = []
    seen: set[str] = set()

    for skill in COMMON_SKILLS:
        if skill in lowered and skill not in seen:
            seen.add(skill)
            keywords.append(skill)

    for match in re.finditer(r"[a-zA-Z][a-zA-Z0-9.+/#-]{1,30}", lowered):
        token = match.group(0).strip(".").strip("-")
        if token in STOPWORDS or len(token) < 3:
            continue
        if token not in seen:
            seen.add(token)
            keywords.append(token)
    return keywords[:25]


def _build_resume_draft(
    jd_text: str,
    resume_text: str,
    analysis: dict[str, Any],
) -> dict[str, Any]:
    lines = [line.strip() for line in resume_text.splitlines() if line.strip()]
    name = lines[0][:60] if lines else "Candidate Name"
    contact = _find_contact_line(lines)
    role_hint = _guess_role_from_jd(jd_text)
    matched = analysis.get("matched_keywords", [])[:8]
    missing = analysis.get("missing_keywords", [])[:5]
    achievements = _extract_resume_bullets(lines)
    if not achievements:
        achievements = [
            "Tailored past experience to align with the target role and highlighted relevant outcomes.",
            "Add a real achievement here with a metric that demonstrates impact.",
            "Showcase one project or responsibility that maps directly to the JD requirements.",
        ]

    summary_bits = matched[:4] or ["relevant experience", "execution", "problem solving"]
    summary = (
        f"Results-focused candidate targeting {role_hint}. "
        f"Highlights experience across {', '.join(summary_bits)} while keeping the resume aligned to the role."
    )

    skills = matched + [f"Validate and add if true: {item}" for item in missing]
    education = _extract_section_lines(lines, ("education", "university", "college", "degree"), 3)
    projects = _extract_section_lines(lines, ("project", "portfolio", "case study"), 3)
    certifications = _extract_section_lines(lines, ("certification", "certificate", "course"), 3)
    notes = [
        "Replace any validation notes with real, provable skills before sharing the resume.",
        "Update placeholders with exact metrics from your own work.",
    ]

    return {
        "name": name,
        "headline": f"{role_hint} Resume",
        "contact": contact,
        "summary": summary,
        "skills": skills[:10],
        "experience": achievements[:6],
        "projects": projects,
        "education": education,
        "certifications": certifications,
        "notes": notes,
    }


def _find_contact_line(lines: list[str]) -> str:
    for line in lines[:5]:
        if "@" in line or re.search(r"\+?\d[\d\s-]{7,}", line):
            return line[:100]
    return "Add email | Add phone | Add LinkedIn"


def _guess_role_from_jd(jd_text: str) -> str:
    for line in jd_text.splitlines():
        cleaned = line.strip()
        if cleaned and len(cleaned.split()) <= 8:
            return cleaned[:60]
    words = jd_text.split()
    return " ".join(words[:6])[:60] if words else "Target Role"


def _extract_resume_bullets(lines: list[str]) -> list[str]:
    bullets = []
    for line in lines:
        cleaned = line.lstrip("-* ").strip()
        if len(cleaned.split()) < 5:
            continue
        if any(verb in cleaned.lower() for verb in ACTION_VERBS) or re.search(r"\d", cleaned):
            bullets.append(cleaned[:220])
    return bullets[:6]


def _extract_section_lines(lines: list[str], keywords: tuple[str, ...], limit: int) -> list[str]:
    matches = []
    for line in lines:
        lowered = line.lower()
        if any(keyword in lowered for keyword in keywords):
            matches.append(line[:180])
    return matches[:limit]


def _generate_resume_files(chat_id: int, analysis: dict[str, Any]) -> list[Path]:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = generated_output_dir("results", chat_id) / timestamp
    output_dir.mkdir(parents=True, exist_ok=True)

    draft = analysis["draft_resume"]
    context = {
        "analysis": analysis,
        "name": draft["name"],
        "headline": draft["headline"],
        "contact": draft["contact"],
        "summary": draft["summary"],
        "skills": draft["skills"],
        "experience": draft["experience"],
        "projects": draft["projects"],
        "education": draft["education"],
        "certifications": draft["certifications"],
        "notes": draft["notes"],
        "score": analysis["score"],
        "alignment_label": analysis["alignment_label"],
    }

    files = [
        _write_text_resume(output_dir / "improved_resume.txt", context),
        _write_markdown_resume(output_dir / "improved_resume.md", context),
        _write_docx_resume(output_dir / "improved_resume.docx", context),
    ]

    env = Environment(
        loader=FileSystemLoader(str(TEMPLATES_DIR)),
        autoescape=select_autoescape(["html"]),
    )
    for template_name, output_name in (
        ("resume_classic.html", "resume_classic.html"),
        ("resume_modern.html", "resume_modern.html"),
        ("resume_compact.html", "resume_compact.html"),
    ):
        rendered = env.get_template(template_name).render(**context)
        output_path = output_dir / output_name
        output_path.write_text(rendered, encoding="utf-8")
        files.append(output_path)
        pdf_path = _render_pdf_if_available(output_path)
        if pdf_path is not None:
            files.append(pdf_path)

    return files


def _write_text_resume(path: Path, context: dict[str, Any]) -> Path:
    lines = [
        context["name"],
        context["headline"],
        context["contact"],
        "",
        "SUMMARY",
        context["summary"],
        "",
        "SKILLS",
        *[f"- {item}" for item in context["skills"]],
        "",
        "EXPERIENCE HIGHLIGHTS",
        *[f"- {item}" for item in context["experience"]],
    ]
    if context["projects"]:
        lines.extend(["", "PROJECTS", *[f"- {item}" for item in context["projects"]]])
    if context["education"]:
        lines.extend(["", "EDUCATION", *[f"- {item}" for item in context["education"]]])
    if context["certifications"]:
        lines.extend(
            ["", "CERTIFICATIONS", *[f"- {item}" for item in context["certifications"]]]
        )
    if context["notes"]:
        lines.extend(["", "NOTES", *[f"- {item}" for item in context["notes"]]])
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _write_markdown_resume(path: Path, context: dict[str, Any]) -> Path:
    sections = [
        f"# {context['name']}",
        context["headline"],
        "",
        context["contact"],
        "",
        "## Summary",
        context["summary"],
        "",
        "## Skills",
        *[f"- {item}" for item in context["skills"]],
        "",
        "## Experience Highlights",
        *[f"- {item}" for item in context["experience"]],
    ]
    if context["projects"]:
        sections.extend(["", "## Projects", *[f"- {item}" for item in context["projects"]]])
    if context["education"]:
        sections.extend(["", "## Education", *[f"- {item}" for item in context["education"]]])
    if context["certifications"]:
        sections.extend(
            ["", "## Certifications", *[f"- {item}" for item in context["certifications"]]]
        )
    if context["notes"]:
        sections.extend(["", "## Notes", *[f"- {item}" for item in context["notes"]]])
    path.write_text("\n".join(sections), encoding="utf-8")
    return path


def _write_docx_resume(path: Path, context: dict[str, Any]) -> Path:
    document = Document()
    document.add_heading(context["name"], level=0)
    document.add_paragraph(context["headline"])
    document.add_paragraph(context["contact"])

    _add_docx_section(document, "Summary", [context["summary"]])
    _add_docx_section(document, "Skills", context["skills"])
    _add_docx_section(document, "Experience Highlights", context["experience"])
    if context["projects"]:
        _add_docx_section(document, "Projects", context["projects"])
    if context["education"]:
        _add_docx_section(document, "Education", context["education"])
    if context["certifications"]:
        _add_docx_section(document, "Certifications", context["certifications"])
    if context["notes"]:
        _add_docx_section(document, "Notes", context["notes"])

    document.save(path)
    return path


def _add_docx_section(document: Document, heading: str, items: list[str]) -> None:
    document.add_heading(heading, level=1)
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def _render_pdf_if_available(html_path: Path) -> Path | None:
    if pdfkit is None:
        return None

    executable = shutil.which("wkhtmltopdf")
    if not executable:
        return None

    configuration = pdfkit.configuration(wkhtmltopdf=executable)
    pdf_path = html_path.with_suffix(".pdf")
    pdfkit.from_file(str(html_path), str(pdf_path), configuration=configuration)
    return pdf_path


def _answer_general_query_locally(
    query: str,
    jd_text: str | None,
    resume_text: str | None,
    last_analysis: dict[str, Any] | None,
) -> str:
    lowered = query.lower()
    if "role" in lowered or "apply" in lowered:
        if resume_text:
            keywords = ", ".join(_extract_keywords(resume_text)[:8]) or "your existing skills"
            return (
                "You can likely target adjacent roles based on your current profile. "
                f"From the available resume context, strong areas include {keywords}. "
                "Possible directions are software developer, backend developer, full-stack developer, "
                "automation engineer, QA automation engineer, data analyst, or platform/support engineering, "
                "depending on what is actually true in your experience."
            )
        return (
            "I can help with that. If you share your resume or a short skills summary, "
            "I can suggest nearby roles, skill gaps, and stronger application targets."
        )

    if "improve" in lowered or "resume" in lowered:
        return (
            "A strong resume usually improves by matching the target role more closely, "
            "using measurable achievements, stronger action verbs, and clearer skills alignment. "
            "If you share a JD and resume, I can make the advice much more specific."
        )

    if "jd" in lowered or "job description" in lowered:
        return (
            "A good JD analysis looks for required skills, tools, years of experience, domain language, "
            "and responsibilities so the resume can be tailored to match those signals."
        )

    if last_analysis:
        return (
            f"Based on the last analysis, your latest score was {last_analysis.get('score', 'unknown')} / 10 "
            f"with alignment marked as {last_analysis.get('alignment_label', 'unknown')}. "
            "You can ask about role fit, missing skills, how to improve bullets, or how to target similar jobs."
        )

    return (
        "I can answer general career questions here as well. "
        "Try asking about role fit, resume improvements, job search strategy, interview prep, or what roles suit your profile."
    )
